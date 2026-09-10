"""Build the synthetic English workbook; requires openpyxl."""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

root = Path(__file__).resolve().parents[1]
(root / 'templates').mkdir(exist_ok=True)
w = Workbook()
ins = w.active
ins.title = 'Instructions'
for row in [
    ['Topic', 'Guidance'],
    ['Purpose', 'Three synthetic, manually authored examples. Not an imported customer workbook or execution report.'],
    ['Format', 'One row per step; repeat case metadata. Keep IDs stable and step numbers contiguous.'],
    ['Data', 'Use runtime fixture aliases and synthetic data. Never include credentials or personal data.'],
    ['Review', 'All examples require review before automation. Confirm contracts and expected results.'],
    ['Workflow', 'Use docs/PROMPTS.md: P07 inventory, P08 normalize, P09 generate and execute.'],
    ['Traceability', 'Preserve workbook name, sheet and source rows in normalized cases and test metadata.'],
    ['Units', 'Integer minor units: 100000 means MXN 1000.00.'],
    ['Execution', 'A converted case is not a passed case. Record execution results separately.'],
]:
    ins.append(row)
s = w.create_sheet('TestCases')
s.append(['case_id', 'requirement_id', 'title', 'priority', 'preconditions', 'step_no', 'action', 'test_data_json', 'expected_result', 'layer', 'automation_status'])
cases = [
    ('TR-001', 'REQ-TRANSFER', 'Successful transfer', [
        ('Create and authenticate an owned fixture', '{"balanceMinor":100000,"currency":"MXN"}', 'Authenticated owner has balance 100000.'),
        ('Submit transfer using owned fixture accounts and a fresh key', '{"amountMinor":10000,"currency":"MXN"}', 'HTTP 201; capture the canonical transaction ID.'),
        ('Read canonical balance and transfers; await ledger projection', '{}', 'Balance 90000; exactly one canonical transfer; ledger contains the same ID and amount.'),
    ]),
    ('TR-002', 'REQ-FUNDS', 'Reject insufficient funds', [
        ('Create and authenticate an owned fixture', '{"balanceMinor":100000,"currency":"MXN"}', 'Balance is 100000 with zero transfers.'),
        ('Submit transfer exceeding available funds', '{"amountMinor":110000,"currency":"MXN"}', 'Rejected for insufficient funds; confirm the response contract during review.'),
        ('Read canonical account and transfers', '{}', 'Balance remains 100000 and transfer count remains zero.'),
    ]),
    ('TR-003', 'REQ-REPLAY', 'Sequential idempotent replay', [
        ('Create and authenticate an owned fixture', '{"balanceMinor":100000,"currency":"MXN"}', 'Owned source and beneficiary available; create one case-specific key.'),
        ('Submit transfer then repeat identical request with the same key', '{"amountMinor":10000,"currency":"MXN"}', 'First response 201; replay 200; both have the same canonical transaction ID.'),
        ('Read canonical balance and transfer records', '{}', 'Balance is 90000; exactly one canonical transfer with the captured ID.'),
    ]),
]
for cid, req, title, steps in cases:
    for n, (action, data, expected) in enumerate(steps, 1):
        s.append([cid, req, title, 'P0', 'Healthy local services; fresh owned fixture; clean it up afterward.', n, action, data, expected, 'API', 'REVIEW_REQUIRED'])
for col, values in [('D', 'P0,P1,P2'), ('J', 'UI,API,HYBRID'), ('K', 'REVIEW_REQUIRED,DRAFT,READY,BLOCKED')]:
    dv = DataValidation(type='list', formula1='"' + values + '"')
    dv.showErrorMessage = True
    s.add_data_validation(dv)
    dv.add(f'{col}2:{col}1000')
for ws in w:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='2457E6')
        c.font = Font(color='FFFFFF', bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.fill = PatternFill('solid', fgColor='F4F7FC' if c.row % 2 == 0 else 'FFFFFF')
            c.font = Font(color='14213D')
            c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[row[0].row].height = 78
for col, width in zip('ABCDEFGHIJK', [16, 19, 28, 12, 42, 11, 48, 54, 65, 14, 23]):
    s.column_dimensions[col].width = width
ins.column_dimensions['A'].width = 20
ins.column_dimensions['B'].width = 105
path = root / 'templates/test-cases-template.xlsx'
w.save(path)
r = load_workbook(path)
rows = list(r['TestCases'].values)
assert len(rows) == 10
assert len({row[0] for row in rows[1:]}) == 3
assert all(row[-1] == 'REVIEW_REQUIRED' for row in rows[1:])
assert not any(c.data_type == 'f' for ws in r for row in ws for c in row)
print('Validated workbook: 3 cases, 9 steps, no formulas.')
